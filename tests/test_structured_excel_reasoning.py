from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from agent.service import ExcelReasoningAgent


def build_demo_workbook(path: Path) -> None:
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["Dealer", "Leads", "Appointments", "ConversionRate", "Category"])
    sales.append(["North Motors", 140, 42, "=C2/B2", '=IF(B2>100,"High","Low")'])
    sales.append(["City Auto", 90, 27, "=C3/B3", '=IF(B3>100,"High","Low")'])
    sales.append(["Summit Cars", 180, 63, "=C4/B4", '=IF(B4>100,"High","Low")'])

    staffing = workbook.create_sheet("Staffing")
    staffing.append(["Dealer", "StaffCount", "StaffingCost"])
    staffing.append(["North Motors", 12, 120000])
    staffing.append(["City Auto", 8, 85000])
    staffing.append(["Summit Cars", 15, 150000])

    workbook.save(path)


def test_structured_pipeline_detects_tables_relationships_and_formulas(tmp_path: Path) -> None:
    workbook_path = tmp_path / "demo.xlsx"
    build_demo_workbook(workbook_path)

    agent = ExcelReasoningAgent()
    bundle = agent.load_bundle(str(workbook_path))

    assert len(bundle.parsed_tables.tables) >= 2
    assert len(bundle.schema.tables) >= 2
    assert any(item.column_name == "Category" for item in bundle.formulas.derived_columns)
    assert any(item.column_name == "ConversionRate" for item in bundle.formulas.derived_columns)
    assert any(
        {relation.left_table, relation.right_table}
        == {"Sales_table_1", "Staffing_table_1"}
        for relation in bundle.relationships.relationships
    )


def test_agent_answers_demo_queries_with_execution_backing(tmp_path: Path) -> None:
    workbook_path = tmp_path / "demo.xlsx"
    build_demo_workbook(workbook_path)

    agent = ExcelReasoningAgent()

    top_dealer = agent.answer_query(str(workbook_path), "Top dealer by leads")
    assert top_dealer.verification.valid
    assert "Summit Cars" in top_dealer.answer_text

    join_query = agent.answer_query(str(workbook_path), "Dealer with highest conversion and staffing impact")
    assert join_query.verification.valid
    assert len(join_query.selected_tables) >= 2
    assert "staffing impact" in join_query.answer_text.lower()
    assert join_query.supporting_data[0]["staffing_impact_score"] is not None

    category_query = agent.answer_query(str(workbook_path), "Which dealers are categorized as High?")
    assert category_query.verification.valid
    returned_dealers = {str(row["Dealer"]) for row in category_query.supporting_data}
    assert returned_dealers == {"North Motors", "Summit Cars"}

    average_query = agent.answer_query(str(workbook_path), "Average conversion rate across all dealers")
    assert average_query.verification.valid
    assert "average conversion rate" in average_query.answer_text.lower()
