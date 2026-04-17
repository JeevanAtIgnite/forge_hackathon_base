"""Workbook parser that detects logical tables inside Excel sheets."""

from __future__ import annotations

from collections import Counter
from io import BytesIO
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from excel_parser.models import Column, Formula, Table, WorkbookTables


class RowBlock(BaseModel):
    """Contiguous non-empty row block inside a worksheet."""

    start_row: int
    end_row: int


class ExcelWorkbookParser:
    """Parse an Excel workbook into logical tables.

    The parser looks for contiguous blocks of populated rows, infers a header
    row, and produces row dictionaries plus captured formulas.
    """

    def parse(self, file_path: str) -> WorkbookTables:
        path = Path(file_path)
        formula_wb = load_workbook(path, data_only=False)
        value_wb = load_workbook(path, data_only=True)
        return self._parse_workbook(
            file_name=path.name,
            workbook_title=path.stem.replace("_", " ").replace("-", " ").title(),
            formula_wb=formula_wb,
            value_wb=value_wb,
        )

    def parse_bytes(self, file_name: str, payload: bytes) -> WorkbookTables:
        formula_wb = load_workbook(BytesIO(payload), data_only=False)
        value_wb = load_workbook(BytesIO(payload), data_only=True)
        return self._parse_workbook(
            file_name=file_name,
            workbook_title=Path(file_name).stem.replace("_", " ").replace("-", " ").title(),
            formula_wb=formula_wb,
            value_wb=value_wb,
        )

    def _parse_workbook(self, file_name: str, workbook_title: str, formula_wb: Any, value_wb: Any) -> WorkbookTables:
        tables: list[Table] = []
        notes: list[str] = []

        for sheet_name in formula_wb.sheetnames:
            formula_ws = formula_wb[sheet_name]
            value_ws = value_wb[sheet_name]

            row_blocks = self._detect_row_blocks(value_ws)
            if not row_blocks:
                continue

            for block_index, block in enumerate(row_blocks, start=1):
                table = self._build_table(
                    formula_ws=formula_ws,
                    value_ws=value_ws,
                    sheet_name=sheet_name,
                    block=block,
                    block_index=block_index,
                )
                if table:
                    tables.append(table)
                else:
                    notes.append(f"Skipped low-confidence block on sheet '{sheet_name}' rows {block.start_row}-{block.end_row}.")

        return WorkbookTables(
            file_name=file_name,
            workbook_title=workbook_title,
            sheet_names=list(formula_wb.sheetnames),
            tables=tables,
            notes=notes,
        )

    def _detect_row_blocks(self, worksheet: Any) -> list[RowBlock]:
        populated_rows: list[int] = []
        for row_idx in range(1, worksheet.max_row + 1):
            non_empty = self._non_empty_cells(worksheet, row_idx)
            if non_empty >= 2:
                populated_rows.append(row_idx)

        if not populated_rows:
            return []

        blocks: list[RowBlock] = []
        start = populated_rows[0]
        previous = populated_rows[0]

        for row_idx in populated_rows[1:]:
            if row_idx - previous > 1:
                blocks.append(RowBlock(start_row=start, end_row=previous))
                start = row_idx
            previous = row_idx

        blocks.append(RowBlock(start_row=start, end_row=previous))
        return blocks

    def _build_table(
        self,
        formula_ws: Any,
        value_ws: Any,
        sheet_name: str,
        block: RowBlock,
        block_index: int,
    ) -> Table | None:
        header_row = self._select_header_row(value_ws, block)
        if header_row is None:
            return None

        active_columns = self._detect_active_columns(formula_ws, value_ws, header_row, block)
        if len(active_columns) < 2:
            return None

        headers = self._dedupe_headers(
            [
                self._normalize_header(value_ws.cell(header_row, col_idx).value, fallback=f"{sheet_name}_column_{col_idx}")
                for col_idx in active_columns
            ]
        )
        column_lookup = {
            get_column_letter(col_idx): headers[position]
            for position, col_idx in enumerate(active_columns)
        }

        rows: list[dict[str, Any]] = []
        formulas: list[Formula] = []
        columns: list[Column] = []

        for position, col_idx in enumerate(active_columns):
            sample_values = self._sample_column_values(value_ws, block, header_row, col_idx)
            columns.append(
                Column(
                    name=headers[position],
                    original_name=str(value_ws.cell(header_row, col_idx).value or headers[position]),
                    index=col_idx,
                    excel_letter=get_column_letter(col_idx),
                    sample_values=sample_values,
                )
            )

        for row_idx in range(header_row + 1, block.end_row + 1):
            record: dict[str, Any] = {}
            populated = False

            for position, col_idx in enumerate(active_columns):
                column_name = headers[position]
                formula_cell = formula_ws.cell(row_idx, col_idx)
                value_cell = value_ws.cell(row_idx, col_idx)
                value = value_cell.value

                if isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                    value = self._evaluate_formula(formula_cell.value, record, column_lookup, fallback=value)

                record[column_name] = value

                if value not in (None, ""):
                    populated = True

                if isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                    formulas.append(
                        Formula(
                            cell=formula_cell.coordinate,
                            row_index=row_idx,
                            column_name=column_name,
                            formula=formula_cell.value,
                            cached_value=value,
                            references=self._extract_references(formula_cell.value),
                        )
                    )

            if populated:
                rows.append(record)

        if not rows:
            return None

        table_name = self._make_table_name(sheet_name, headers, block_index)
        return Table(
            name=table_name,
            sheet_name=sheet_name,
            header_row_index=header_row,
            start_row=block.start_row,
            end_row=block.end_row,
            start_column=min(active_columns),
            end_column=max(active_columns),
            columns=columns,
            rows=rows,
            formulas=formulas,
        )

    def _select_header_row(self, worksheet: Any, block: RowBlock) -> int | None:
        candidates = range(block.start_row, min(block.start_row + 3, block.end_row) + 1)
        best_row: int | None = None
        best_score = -1.0

        for row_idx in candidates:
            values = [worksheet.cell(row_idx, col_idx).value for col_idx in range(1, worksheet.max_column + 1)]
            non_empty_values = [value for value in values if value not in (None, "")]
            if len(non_empty_values) < 2:
                continue

            string_like = sum(1 for value in non_empty_values if isinstance(value, str))
            unique_ratio = len({str(value).strip().lower() for value in non_empty_values}) / len(non_empty_values)
            next_row_populated = self._non_empty_cells(worksheet, min(row_idx + 1, worksheet.max_row))

            score = (string_like / len(non_empty_values)) + unique_ratio + (0.2 if next_row_populated >= 2 else 0.0)
            if score > best_score:
                best_score = score
                best_row = row_idx

        return best_row

    def _detect_active_columns(self, formula_ws: Any, value_ws: Any, header_row: int, block: RowBlock) -> list[int]:
        active_columns: list[int] = []
        max_column = max(formula_ws.max_column, value_ws.max_column)
        for col_idx in range(1, max_column + 1):
            header_value = value_ws.cell(header_row, col_idx).value
            if header_value in (None, ""):
                header_value = formula_ws.cell(header_row, col_idx).value

            column_populated = any(
                value_ws.cell(row_idx, col_idx).value not in (None, "")
                for row_idx in range(header_row + 1, block.end_row + 1)
            )
            has_formula = any(
                isinstance(formula_ws.cell(row_idx, col_idx).value, str)
                and formula_ws.cell(row_idx, col_idx).value.startswith("=")
                for row_idx in range(header_row + 1, block.end_row + 1)
            )
            if header_value not in (None, "") and (column_populated or has_formula):
                active_columns.append(col_idx)

        return active_columns

    def _sample_column_values(self, worksheet: Any, block: RowBlock, header_row: int, col_idx: int) -> list[Any]:
        values: list[Any] = []
        for row_idx in range(header_row + 1, block.end_row + 1):
            value = worksheet.cell(row_idx, col_idx).value
            if value not in (None, ""):
                values.append(value)
            if len(values) == 3:
                break
        return values

    def _non_empty_cells(self, worksheet: Any, row_idx: int) -> int:
        count = 0
        for col_idx in range(1, worksheet.max_column + 1):
            value = worksheet.cell(row_idx, col_idx).value
            if value not in (None, ""):
                count += 1
        return count

    def _normalize_header(self, value: Any, fallback: str) -> str:
        text = str(value or fallback).strip()
        text = re.sub(r"\s+", " ", text)
        text = re.sub(r"[^a-zA-Z0-9_ ]", "", text)
        text = text.replace(" ", "_")
        text = re.sub(r"_+", "_", text).strip("_")
        return text or fallback

    def _dedupe_headers(self, headers: list[str]) -> list[str]:
        counts: Counter[str] = Counter()
        unique_headers: list[str] = []
        for header in headers:
            counts[header] += 1
            if counts[header] == 1:
                unique_headers.append(header)
            else:
                unique_headers.append(f"{header}_{counts[header]}")
        return unique_headers

    def _extract_references(self, formula: str) -> list[str]:
        references = re.findall(r"\b[A-Z]{1,3}\d+\b", formula)
        return sorted(set(references))

    def _evaluate_formula(
        self,
        formula: str,
        record: dict[str, Any],
        column_lookup: dict[str, str],
        fallback: Any = None,
    ) -> Any:
        expression = formula.strip()
        if not expression.startswith("="):
            return fallback

        body = expression[1:].strip()
        if body.upper().startswith("IF(") and body.endswith(")"):
            return self._evaluate_if(body[3:-1], record, column_lookup, fallback)

        return self._evaluate_expression(body, record, column_lookup, fallback)

    def _evaluate_if(
        self,
        args_expression: str,
        record: dict[str, Any],
        column_lookup: dict[str, str],
        fallback: Any,
    ) -> Any:
        parts = self._split_function_args(args_expression)
        if len(parts) != 3:
            return fallback

        condition, true_case, false_case = parts
        condition_result = self._evaluate_condition(condition, record, column_lookup)
        branch = true_case if condition_result else false_case
        return self._evaluate_literal_or_expression(branch, record, column_lookup, fallback)

    def _evaluate_condition(self, condition: str, record: dict[str, Any], column_lookup: dict[str, str]) -> bool:
        translated = self._translate_expression(condition, record, column_lookup)
        translated = translated.replace("<>", "!=")
        translated = re.sub(r"(?<![<>=!])=(?!=)", "==", translated)

        try:
            return bool(eval(translated, {"__builtins__": {}}, {}))
        except Exception:
            return False

    def _evaluate_literal_or_expression(
        self,
        value: str,
        record: dict[str, Any],
        column_lookup: dict[str, str],
        fallback: Any,
    ) -> Any:
        cleaned = value.strip()
        if cleaned.startswith('"') and cleaned.endswith('"'):
            return cleaned[1:-1]
        if re.fullmatch(r"-?\d+(\.\d+)?", cleaned):
            return float(cleaned) if "." in cleaned else int(cleaned)
        return self._evaluate_expression(cleaned, record, column_lookup, fallback)

    def _evaluate_expression(
        self,
        expression: str,
        record: dict[str, Any],
        column_lookup: dict[str, str],
        fallback: Any,
    ) -> Any:
        translated = self._translate_expression(expression, record, column_lookup)
        if re.search(r"[A-Za-z_]", translated):
            return fallback

        try:
            return eval(translated, {"__builtins__": {}}, {})
        except Exception:
            return fallback

    def _translate_expression(self, expression: str, record: dict[str, Any], column_lookup: dict[str, str]) -> str:
        def replace(match: re.Match[str]) -> str:
            reference = match.group(0)
            letter_match = re.match(r"([A-Z]{1,3})\d+", reference)
            if not letter_match:
                return "None"

            column_name = column_lookup.get(letter_match.group(1))
            value = record.get(column_name)
            return repr(value)

        return re.sub(r"\b[A-Z]{1,3}\d+\b", replace, expression)

    def _split_function_args(self, expression: str) -> list[str]:
        parts: list[str] = []
        current: list[str] = []
        depth = 0
        in_quotes = False

        for char in expression:
            if char == '"':
                in_quotes = not in_quotes
            elif not in_quotes:
                if char == "(":
                    depth += 1
                elif char == ")":
                    depth = max(0, depth - 1)
                elif char == "," and depth == 0:
                    parts.append("".join(current).strip())
                    current = []
                    continue
            current.append(char)

        if current:
            parts.append("".join(current).strip())
        return parts

    def _make_table_name(self, sheet_name: str, headers: list[str], block_index: int) -> str:
        return f"{sheet_name}_table_{block_index}"
